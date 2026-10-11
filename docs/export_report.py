"""Export the source report to printable PDF and editable Word documents."""
from pathlib import Path
from html import escape
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Preformatted,
)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "generated-artifacts" / "documentation"
OUT.mkdir(parents=True, exist_ok=True)
SOURCE = ROOT / "docs" / "PROJECT_REPORT.md"
ORANGE = "#C2410C"
styles = getSampleStyleSheet()
styles.add(ParagraphStyle("ReportBody", fontName="Helvetica", fontSize=9.5,
                          leading=13, textColor=colors.HexColor("#334155"), spaceAfter=6))
styles.add(ParagraphStyle("ReportTitle", fontName="Helvetica-Bold", fontSize=30,
                          leading=36, alignment=TA_CENTER, textColor=colors.HexColor(ORANGE)))
styles.add(ParagraphStyle("ReportSubtitle", fontName="Helvetica", fontSize=16,
                          leading=24, alignment=TA_CENTER, textColor=colors.HexColor("#475569")))
styles.add(ParagraphStyle("ReportChapter", fontName="Helvetica-Bold", fontSize=19,
                          leading=24, spaceAfter=18, textColor=colors.HexColor(ORANGE)))
styles.add(ParagraphStyle("ReportSection", fontName="Helvetica-Bold", fontSize=12,
                          leading=17, spaceBefore=10, spaceAfter=7, textColor=colors.HexColor("#0F172A")))
styles.add(ParagraphStyle("ReportCell", parent=styles["ReportBody"], fontSize=8.5,
                          leading=11, spaceAfter=0))
styles.add(ParagraphStyle("ReportCover", parent=styles["ReportBody"], alignment=TA_CENTER,
                          spaceAfter=15))
styles.add(ParagraphStyle("ReportCode", fontName="Courier", fontSize=8,
                          leading=11, textColor=colors.HexColor("#334155")))

document = Document()
section = document.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.7)
normal = document.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(10.5)
normal.paragraph_format.space_after = Pt(8)
for name in ["Title", "Heading 1", "Heading 2"]:
    document.styles[name].font.color.rgb = RGBColor.from_string("C2410C")
document.core_properties.title = "Indian Dumb Charades - Project Submission Report"
document.core_properties.subject = "Project features, architecture, installation, testing and limitations"
document.core_properties.author = "Project Documentation"
footer = section.footer.paragraphs[0]
footer.text = "Indian Dumb Charades | Project Submission Report"
footer.style = document.styles["Caption"]
story = []
cover = True
lines = SOURCE.read_text(encoding="utf-8").splitlines()

def add_paragraph(text, style="ReportBody"):
    story.append(Paragraph(escape(text), styles[style]))
    p = document.add_paragraph(text)
    if style == "ReportCover":
        p.alignment = 1
    return p

def add_table(rows):
    widths = [170, 325] if len(rows[0]) == 2 else [95, 160, 240]
    cells = [[Paragraph(escape(cell), styles["ReportCell"]) for cell in row] for row in rows]
    table = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FFEDD5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#FDBA74")),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.extend([table, Spacer(1, 12)])
    word_table = document.add_table(rows=0, cols=len(rows[0]))
    word_table.style = "Light Shading Accent 1"
    for row in rows:
        for cell, text in zip(word_table.add_row().cells, row):
            cell.text = text
    document.add_paragraph()

i = 0
while i < len(lines):
    line = lines[i].strip()
    if not line:
        i += 1
        continue
    if line.startswith("# "):
        story.extend([Spacer(1, 105), Paragraph(escape(line[2:]), styles["ReportTitle"]), Spacer(1, 18)])
        document.add_paragraph(line[2:], "Title")
    elif line.startswith("## "):
        cover = False
        story.append(PageBreak())
        document.add_page_break()
        story.append(Paragraph(escape(line[3:]), styles["ReportChapter"]))
        document.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        story.append(Paragraph(escape(line[4:]), styles["ReportSection"]))
        document.add_heading(line[4:], level=2)
    elif line.startswith("```"):
        code = []
        i += 1
        while i < len(lines) and not lines[i].startswith("```"):
            code.append(lines[i])
            i += 1
        story.extend([Preformatted("\n".join(code), styles["ReportCode"], maxLineLength=94), Spacer(1, 10)])
        p = document.add_paragraph("\n".join(code))
        for run in p.runs:
            run.font.name, run.font.size = "Consolas", Pt(8)
    elif line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                rows.append(cells)
            i += 1
        add_table(rows)
        continue
    elif line.startswith("- "):
        story.append(Paragraph("&#8226; " + escape(line[2:]), styles["ReportBody"]))
        document.add_paragraph(line[2:], "List Bullet")
    elif re.match(r"^\d+\. ", line):
        story.append(Paragraph(escape(line), styles["ReportBody"]))
        document.add_paragraph(re.sub(r"^\d+\. ", "", line), "List Number")
    else:
        if cover and line == "Project Submission Report":
            story.extend([Paragraph(line, styles["ReportSubtitle"]), Spacer(1, 65)])
            document.add_paragraph(line, "Subtitle")
        else:
            add_paragraph(line, "ReportCover" if cover else "ReportBody")
    i += 1

def decorate(canvas, doc):
    canvas.saveState()
    w, h = doc.pagesize
    if doc.page > 1:
        canvas.setStrokeColor(colors.HexColor("#E2E8F0"))
        canvas.line(50, h - 36, w - 50, h - 36)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(50, h - 27, "INDIAN DUMB CHARADES")
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#64748B"))
    canvas.drawString(50, 27, "Project Submission Report | 11 October 2026")
    canvas.drawRightString(w - 50, 27, str(doc.page))
    canvas.restoreState()

pdf_path = OUT / "Indian_Dumb_Charades_Project_Report.pdf"
docx_path = OUT / "Indian_Dumb_Charades_Project_Report.docx"
SimpleDocTemplate(
    str(pdf_path), pagesize=(595.28, 841.89), rightMargin=50, leftMargin=50,
    topMargin=52, bottomMargin=48, title="Indian Dumb Charades - Project Submission Report",
    author="Project Documentation",
).build(story, onFirstPage=decorate, onLaterPages=decorate)
document.save(docx_path)
print(f"Created {pdf_path.relative_to(ROOT)}")
print(f"Created {docx_path.relative_to(ROOT)}")
